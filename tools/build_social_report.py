import argparse
import calendar
import json
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path
from urllib.parse import unquote

import duckdb
from openpyxl import load_workbook


def parse_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return datetime.strptime(value.strip(), "%d/%m/%Y").date()
        except ValueError:
            return None
    return None


def integer(value: object) -> int:
    if value in (None, ""):
        return 0
    return int(float(str(value).replace(",", ".")))


def normalized_text(value: str) -> str:
    decomposed = unicodedata.normalize("NFKD", value)
    return "".join(character for character in decomposed if not unicodedata.combining(character)).casefold().strip()


def parse_percentage(value: object) -> tuple[float, str] | None:
    if value is None:
        return None
    label = str(value).replace("\xa0", " ").strip()
    match = re.search(r"(\d+(?:[.,]\d+)?)", label)
    if not match:
        return None
    amount = float(match.group(1).replace(",", "."))
    is_less_than = "<" in label or "menos" in normalized_text(label)
    plotted_value = amount / 2 if is_less_than else amount
    display_label = f"<{amount:g}%" if is_less_than else f"{amount:g}%"
    return plotted_value, display_label


def read_audience(worksheet):
    category_ids = {
        "empresa": "companies",
        "ubicacion": "locations",
        "tamano de la empresa": "company_size",
        "nivel de responsabilidad": "seniority",
        "cargo": "job_titles",
        "sector": "industries",
    }
    category_labels = {
        "companies": "Companies",
        "locations": "Locations",
        "company_size": "Company size",
        "seniority": "Seniority",
        "job_titles": "Job titles",
        "industries": "Industries",
    }
    city_coordinates = {
        "buenos aires y alrededores": (-58.3816, -34.6037),
        "area metropolitana de lima": (-77.0428, -12.0464),
        "area metropolitana de bogota d.c.": (-74.0721, 4.7110),
        "area metropolitana de bogota": (-74.0721, 4.7110),
        "area metropolitana de santiago": (-70.6693, -33.4489),
        "area metropolitana de ciudad de mexico": (-99.1332, 19.4326),
        "cordoba y alrededores, argentina": (-64.1888, -31.4201),
        "madrid y alrededores": (-3.7038, 40.4168),
        "sao paulo y alrededores": (-46.6333, -23.5505),
        "area metropolitana de medellin": (-75.5812, 6.2442),
        "barcelona y alrededores": (2.1734, 41.3851),
        "bangalore y alrededores": (77.5946, 12.9716),
        "paris y alrededores": (2.3522, 48.8566),
    }
    groups = {}
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if len(row) < 3 or row[0] is None or row[1] is None:
            continue
        category = category_ids.get(normalized_text(str(row[0])))
        percentage = parse_percentage(row[2])
        if category is None or percentage is None:
            continue
        plotted_value, display_label = percentage
        item = {
            "name": str(row[1]).strip(),
            "value": plotted_value,
            "label": display_label,
        }
        if category == "locations":
            coordinates = city_coordinates.get(normalized_text(item["name"]))
            if coordinates:
                item["longitude"], item["latitude"] = coordinates
        groups.setdefault(category, []).append(item)
    return {
        key: {
            "label": category_labels[key],
            "items": sorted(items, key=lambda item: item["value"], reverse=True),
        }
        for key, items in groups.items()
    }


def metric_rows(worksheet, date_column: int, value_columns: tuple[int, ...]):
    rows = []
    for values in worksheet.iter_rows(values_only=True):
        parsed_date = parse_date(values[date_column] if len(values) > date_column else None)
        if parsed_date is None:
            continue
        rows.append(
            (
                parsed_date,
                *(integer(values[index] if len(values) > index else 0) for index in value_columns),
            )
        )
    return rows


def find_sheet(workbook, prefix: str):
    for worksheet in workbook.worksheets:
        if worksheet.title.startswith(prefix):
            return worksheet
    raise ValueError(f"Required worksheet not found: {prefix}")


def read_posts(worksheet, url_column: int, date_column: int, metric_column: int):
    posts = []
    for values in worksheet.iter_rows(min_row=4, values_only=True):
        if len(values) <= max(url_column, date_column, metric_column):
            continue
        url = values[url_column]
        post_date = parse_date(values[date_column])
        if not isinstance(url, str) or not url.startswith("https://www.linkedin.com/") or post_date is None:
            continue
        posts.append(
            {
                "date": post_date.isoformat(),
                "url": url,
                "value": integer(values[metric_column]),
                "title": post_title(url),
            }
        )
    return posts


def post_title(url: str) -> str:
    slug = url.split("_", 1)[-1].split("-ugcPost-", 1)[0]
    slug = re.sub(r"-\d{15,}.*$", "", slug)
    words = unquote(slug).replace("-", " ").split()
    if words and words[0].lower() == "im":
        words[0] = "I'm"
    elif words and words[0].lower() == "its":
        words[0] = "It's"
    title = " ".join(words[:9]).strip()
    return title[:1].upper() + title[1:] if title else "LinkedIn post"


def main():
    parser = argparse.ArgumentParser(description="Build the public static LinkedIn report data.")
    parser.add_argument("--workbook", required=True, type=Path, help="Path to the LinkedIn analytics Excel export.")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "social-media" / "data" / "report.json",
        help="Output path for the dashboard JSON.",
    )
    args = parser.parse_args()
    if not args.workbook.is_file():
        parser.error(f"Workbook does not exist: {args.workbook}")

    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    summary_sheet = find_sheet(workbook, "DESCUBRIMIENTO")
    interaction_rows = metric_rows(find_sheet(workbook, "INTERACCI"), 0, (1, 2))
    follower_sheet = find_sheet(workbook, "SEGUID")
    follower_rows = metric_rows(follower_sheet, 0, (1,))
    post_sheet = find_sheet(workbook, "PUBLICACIONES")
    audience_sheets = [sheet for sheet in workbook.worksheets if sheet.title.startswith("INFORMACI")]
    if len(audience_sheets) < 2:
        raise ValueError("The workbook must contain both LinkedIn audience detail sheets.")
    follower_audience = read_audience(audience_sheets[0])
    content_audience = read_audience(audience_sheets[1])
    top_engagement = sorted(read_posts(post_sheet, 0, 1, 2), key=lambda post: post["value"], reverse=True)[:10]
    top_impressions = sorted(read_posts(post_sheet, 4, 5, 6), key=lambda post: post["value"], reverse=True)[:10]

    if not interaction_rows or not follower_rows:
        raise ValueError("The workbook does not contain daily interaction and follower data.")

    summary = {
        str(row[0]).strip(): integer(row[1])
        for row in summary_sheet.iter_rows(values_only=True)
        if len(row) > 1 and row[0] in ("Impresiones", "Miembros alcanzados")
    }
    follower_total = next(
        (
            integer(row[1])
            for row in follower_sheet.iter_rows(min_row=1, max_row=1, values_only=True)
            if len(row) > 1
        ),
        0,
    )

    connection = duckdb.connect(":memory:")
    try:
        connection.execute(
            "CREATE TABLE daily (day DATE, impressions BIGINT, engagements BIGINT, new_followers BIGINT)"
        )
        follower_by_date = {row[0]: row[1] for row in follower_rows}
        connection.executemany(
            "INSERT INTO daily VALUES (?, ?, ?, ?)",
            [
                (day, impressions, engagements, follower_by_date.get(day, 0))
                for day, impressions, engagements in interaction_rows
            ],
        )
        daily = connection.execute(
            """
            SELECT day, impressions, engagements, new_followers
            FROM daily
            ORDER BY day
            """
        ).fetchall()
        monthly = connection.execute(
            """
            SELECT
                date_trunc('month', day)::DATE AS month_start,
                count(*)::INTEGER AS days_recorded,
                sum(impressions)::BIGINT AS impressions,
                sum(engagements)::BIGINT AS engagements,
                sum(new_followers)::BIGINT AS new_followers
            FROM daily
            GROUP BY 1
            ORDER BY 1
            """
        ).fetchall()
    finally:
        connection.close()
        workbook.close()

    impressions = sum(row[1] for row in daily)
    engagements = sum(row[2] for row in daily)
    report = {
        "period": {
            "start": daily[0][0].isoformat(),
            "end": daily[-1][0].isoformat(),
        },
        "summary": {
            "impressions": summary.get("Impresiones", impressions),
            "reach": summary.get("Miembros alcanzados", 0),
            "engagements": engagements,
            "engagement_rate": round(engagements / impressions * 100, 2) if impressions else 0,
            "followers": follower_total,
            "new_followers": sum(row[3] for row in daily),
            "estimated_start_followers": max(follower_total - sum(row[3] for row in daily), 0),
        },
        "audiences": {
            "followers": follower_audience,
            "content": content_audience,
        },
        "daily": [
            {
                "date": row[0].isoformat(),
                "impressions": row[1],
                "engagements": row[2],
                "new_followers": row[3],
            }
            for row in daily
        ],
        "monthly": [
            {
                "month": row[0].isoformat(),
                "days_recorded": row[1],
                "impressions": row[2],
                "engagements": row[3],
                "new_followers": row[4],
            }
            for row in monthly
        ],
        "top_engagement": top_engagement,
        "top_impressions": top_impressions,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    print(f"Built {args.output} with {len(daily)} daily records and {len(monthly)} monthly records.")
    from build_bokeh_dashboard import build_dashboard

    map_path = args.output.parent / "world-land.geojson"
    if not map_path.is_file():
        raise FileNotFoundError(f"Required world map asset not found: {map_path}")
    build_dashboard(args.output, map_path, args.output.parent / "bokeh-dashboard.json")


if __name__ == "__main__":
    main()
